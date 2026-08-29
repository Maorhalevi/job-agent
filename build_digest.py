from openpyxl import load_workbook
import datetime as dt, html

wb = load_workbook('jobs.xlsx')
ws = wb.active
rows = []
for r in ws.iter_rows(min_row=2):
    company, role, loc, pct, why = r[0].value, r[1].value, r[2].value, r[3].value, r[4].value
    link = r[5].hyperlink.target if r[5].hyperlink else (r[5].value or '')
    if company and 'No strong matches' in str(company):
        continue
    rows.append((company, role, loc, pct, why, link))

today = dt.date.today().strftime('%A, %B %-d, %Y')
esc = html.escape

def pct_str(p):
    try: return f"{round(float(p)*100)}%"
    except: return str(p)

items = ""
for company, role, loc, pct, why, link in rows:
    items += f"""
      <tr>
        <td style="padding:12px 14px;border-bottom:1px solid #e5e7eb;vertical-align:top;">
          <div style="font-weight:600;color:#111827;font-size:15px;">{esc(str(role))}</div>
          <div style="color:#6b7280;font-size:13px;margin-top:2px;">{esc(str(company))} &middot; {esc(str(loc))}</div>
          <div style="color:#6b7280;font-size:12px;margin-top:4px;">{esc(str(why))}</div>
        </td>
        <td style="padding:12px 14px;border-bottom:1px solid #e5e7eb;vertical-align:top;text-align:center;white-space:nowrap;">
          <span style="display:inline-block;background:#16263F;color:#fff;border-radius:12px;padding:3px 10px;font-size:13px;font-weight:600;">{pct_str(pct)}</span>
        </td>
        <td style="padding:12px 14px;border-bottom:1px solid #e5e7eb;vertical-align:top;text-align:right;white-space:nowrap;">
          <a href="{esc(str(link))}" style="color:#1A4FD6;text-decoration:none;font-weight:600;font-size:14px;">Apply &rarr;</a>
        </td>
      </tr>"""

digest = f"""<!doctype html>
<html><body style="margin:0;background:#f3f4f6;font-family:Arial,Helvetica,sans-serif;">
  <div style="max-width:640px;margin:0 auto;padding:24px 16px;">
    <div style="background:#16263F;border-radius:10px 10px 0 0;padding:20px 24px;">
      <div style="color:#fff;font-size:20px;font-weight:700;">Daily PM Jobs</div>
      <div style="color:#9fb0c9;font-size:13px;margin-top:4px;">Product Manager &middot; Cybersecurity &middot; Israel &mdash; {today}</div>
    </div>
    <div style="background:#fff;border-radius:0 0 10px 10px;padding:8px 10px 4px;">
      <table style="width:100%;border-collapse:collapse;">{items}
      </table>
    </div>
    <div style="color:#9ca3af;font-size:12px;text-align:center;margin-top:16px;">
      {len(rows)} matched role(s) with a &ge;60% fit. Generated automatically by your job agent.
    </div>
  </div>
</body></html>"""

with open('digest.html', 'w') as f:
    f.write(digest)
print(f"Wrote digest.html with {len(rows)} roles")
