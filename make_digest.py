#!/usr/bin/env python3
"""Render the matched jobs from jobs.xlsx into an email-friendly digest.html.

The routine emails the HTML in digest.html. The scraping/scoring script
(job_agent.py) writes jobs.xlsx; this converts that output into HTML with
inline styles (email clients strip <style>/external CSS), preserving the
real apply links behind the "Open" hyperlinks.
"""
import datetime as dt
import html
from openpyxl import load_workbook

wb = load_workbook("jobs.xlsx")
ws = wb.active

rows = []
for r in range(2, ws.max_row + 1):
    company = ws.cell(r, 1).value
    role = ws.cell(r, 2).value
    location = ws.cell(r, 3).value
    pct = ws.cell(r, 4).value
    why = ws.cell(r, 5).value
    link_cell = ws.cell(r, 6)
    url = link_cell.hyperlink.target if link_cell.hyperlink else ""
    if not role:  # the "no matches" placeholder row has an empty role/link
        continue
    rows.append({
        "company": company or "",
        "role": role or "",
        "location": location or "",
        "pct": int(round((pct or 0) * 100)),
        "why": why or "",
        "url": url or "",
    })

today = dt.date.today().strftime("%A, %d %B %Y")


def esc(s):
    return html.escape(str(s))


cards = []
for j in rows:
    badge = "#1b7f3b" if j["pct"] >= 80 else ("#b8860b" if j["pct"] >= 70 else "#5a6b7b")
    link_html = (
        f'<a href="{esc(j["url"])}" style="display:inline-block;margin-top:10px;'
        f'padding:8px 16px;background:#16263F;color:#ffffff;text-decoration:none;'
        f'border-radius:6px;font-size:13px;font-weight:bold;">Apply / View role &rarr;</a>'
        if j["url"] else ""
    )
    cards.append(f"""
      <tr>
        <td style="padding:16px 18px;border:1px solid #e3e8ee;border-radius:10px;">
          <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse;">
            <tr>
              <td style="font-size:17px;font-weight:bold;color:#16263F;">{esc(j['role'])}</td>
              <td align="right" style="white-space:nowrap;">
                <span style="display:inline-block;padding:4px 10px;background:{badge};color:#ffffff;border-radius:12px;font-size:13px;font-weight:bold;">{j['pct']}% match</span>
              </td>
            </tr>
          </table>
          <div style="margin-top:4px;font-size:14px;color:#334;">{esc(j['company'])}</div>
          <div style="margin-top:2px;font-size:13px;color:#6b7684;">&#128205; {esc(j['location'])}</div>
          <div style="margin-top:8px;font-size:13px;color:#445;">Why it fits: {esc(j['why'])}</div>
          {link_html}
        </td>
      </tr>
      <tr><td style="height:12px;line-height:12px;">&nbsp;</td></tr>""")

cards_html = "".join(cards)

digest = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Daily PM jobs</title>
</head>
<body style="margin:0;padding:0;background:#f4f6f9;font-family:Arial,Helvetica,sans-serif;">
  <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#f4f6f9;padding:24px 0;">
    <tr>
      <td align="center">
        <table role="presentation" width="640" cellpadding="0" cellspacing="0" style="max-width:640px;width:100%;background:#ffffff;border-radius:14px;overflow:hidden;border:1px solid #e3e8ee;">
          <tr>
            <td style="background:#16263F;padding:24px 24px 20px 24px;">
              <div style="font-size:22px;font-weight:bold;color:#ffffff;">Daily PM jobs</div>
              <div style="margin-top:4px;font-size:13px;color:#aebcd0;">Matched Product Manager roles &middot; {today}</div>
            </td>
          </tr>
          <tr>
            <td style="padding:22px 24px 8px 24px;">
              <div style="font-size:14px;color:#334;margin-bottom:16px;">
                Found <strong>{len(rows)}</strong> role{'s' if len(rows) != 1 else ''} matching your profile today
                (cybersecurity / cloud Product Management, Israel &amp; remote).
              </div>
              <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-collapse:separate;">
                {cards_html}
              </table>
            </td>
          </tr>
          <tr>
            <td style="padding:8px 24px 24px 24px;border-top:1px solid #eef1f5;">
              <div style="font-size:11px;color:#9aa6b2;line-height:1.5;">
                Sources: Greenhouse / Lever / Ashby ATS feeds for curated cyber &amp; cloud companies.
                Match scores are keyword-based this run (no AI scoring key configured).
                Generated automatically by your Daily PM Job Hunt routine.
              </div>
            </td>
          </tr>
        </table>
      </td>
    </tr>
  </table>
</body>
</html>"""

with open("digest.html", "w", encoding="utf-8") as f:
    f.write(digest)

print(f"Wrote digest.html with {len(rows)} job(s)")
